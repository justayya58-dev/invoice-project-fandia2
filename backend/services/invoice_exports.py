import base64
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from models.auditpro import CompanyProfile, Customer, Invoice


ONES = ["nol", "satu", "dua", "tiga", "empat", "lima", "enam", "tujuh", "delapan", "sembilan", "sepuluh", "sebelas"]


def terbilang(value: int) -> str:
    value = int(round(value))
    if value < 0:
        return "minus " + terbilang(abs(value))
    if value < 12:
        return ONES[value]
    if value < 20:
        return terbilang(value - 10) + " belas"
    if value < 100:
        return terbilang(value // 10) + " puluh" + (" " + terbilang(value % 10) if value % 10 else "")
    if value < 200:
        return "seratus" + (" " + terbilang(value - 100) if value > 100 else "")
    if value < 1000:
        return terbilang(value // 100) + " ratus" + (" " + terbilang(value % 100) if value % 100 else "")
    if value < 2000:
        return "seribu" + (" " + terbilang(value - 1000) if value > 1000 else "")
    if value < 1_000_000:
        return terbilang(value // 1000) + " ribu" + (" " + terbilang(value % 1000) if value % 1000 else "")
    if value < 1_000_000_000:
        return terbilang(value // 1_000_000) + " juta" + (" " + terbilang(value % 1_000_000) if value % 1_000_000 else "")
    if value < 1_000_000_000_000:
        return terbilang(value // 1_000_000_000) + " miliar" + (" " + terbilang(value % 1_000_000_000) if value % 1_000_000_000 else "")
    return terbilang(value // 1_000_000_000_000) + " triliun" + (" " + terbilang(value % 1_000_000_000_000) if value % 1_000_000_000_000 else "")


def amount_words(value: float) -> str:
    words = terbilang(int(round(value))).strip()
    return (words[:1].upper() + words[1:] + " rupiah") if words else "Nol rupiah"


def _pdf_escape(value: str) -> str:
    return value.encode("latin-1", "replace").decode("latin-1").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _rgb(hex_color: str) -> tuple[float, float, float]:
    clean = hex_color.lstrip("#")
    try:
        return tuple(int(clean[index:index + 2], 16) / 255 for index in (0, 2, 4))  # type: ignore[return-value]
    except (ValueError, TypeError):
        return (0.06, 0.72, 0.51)


def _wrap(text: str, limit: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > limit and current:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def _jpeg_info(data_url: str | None) -> tuple[bytes, int, int] | None:
    if not data_url or not data_url.startswith("data:image/jpeg;base64,"):
        return None
    try:
        data = base64.b64decode(data_url.split(",", 1)[1])
        index = 2
        while index < len(data) - 9:
            if data[index] != 0xFF:
                index += 1
                continue
            marker = data[index + 1]
            if marker in {0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF}:
                height = int.from_bytes(data[index + 5:index + 7], "big")
                width = int.from_bytes(data[index + 7:index + 9], "big")
                return data, width, height
            length = int.from_bytes(data[index + 2:index + 4], "big")
            index += 2 + max(length, 2)
    except Exception:
        return None
    return None


def build_invoice_pdf(invoice: Invoice, company: CompanyProfile, customer: Customer | None) -> bytes:
    accent = _rgb(company.accent_color)
    if invoice.template == "classic":
        accent = (0.18, 0.22, 0.28)
    elif invoice.template == "minimal":
        accent = (0.12, 0.12, 0.12)
    commands: list[str] = []

    def text(x: float, y: float, size: int, value: str, bold: bool = False, color: tuple[float, float, float] = (0.1, 0.13, 0.18)) -> None:
        font = "F2" if bold else "F1"
        commands.append(f"BT /{font} {size} Tf {color[0]:.3f} {color[1]:.3f} {color[2]:.3f} rg {x} {y} Td ({_pdf_escape(value)}) Tj ET")

    if invoice.template == "modern":
        commands.append(f"{accent[0]:.3f} {accent[1]:.3f} {accent[2]:.3f} rg 0 742 612 50 re f")
    elif invoice.template == "classic":
        commands.append("0.75 0.75 0.75 RG 28 28 556 736 re S")
    else:
        commands.append(f"{accent[0]:.3f} {accent[1]:.3f} {accent[2]:.3f} RG 40 735 m 572 735 l S")

    logo = _jpeg_info(company.logo_data_url)
    if logo:
        commands.append("q 72 0 0 44 40 700 cm /Im1 Do Q")
        company_x = 125
    else:
        commands.append(f"{accent[0]:.3f} {accent[1]:.3f} {accent[2]:.3f} rg 40 696 48 48 re f")
        text(55, 712, 16, "".join(part[0] for part in company.name.split()[:2]).upper(), True, (1, 1, 1))
        company_x = 100
    header_color = (1, 1, 1) if invoice.template == "modern" else (0.1, 0.13, 0.18)
    text(company_x, 720, 16, company.name, True, header_color)
    text(company_x, 704, 8, company.email + " | " + company.phone, False, header_color)
    text(440, 716, 24, "INVOICE", True, header_color)
    text(442, 700, 9, invoice.number, False, header_color)
    text(40, 665, 8, "DARI", True, accent)
    text(40, 648, 10, company.name, True)
    for index, line in enumerate(_wrap(company.address, 48)[:2]):
        text(40, 633 - index * 12, 8, line, False, (0.35, 0.38, 0.43))
    text(320, 665, 8, "TAGIH KEPADA", True, accent)
    text(320, 648, 10, invoice.customer, True)
    if customer:
        text(320, 633, 8, customer.email or customer.phone, False, (0.35, 0.38, 0.43))
        for index, line in enumerate(_wrap(customer.address, 42)[:2]):
            text(320, 620 - index * 12, 8, line, False, (0.35, 0.38, 0.43))
    text(40, 585, 8, "TANGGAL", True, (0.35, 0.38, 0.43))
    text(40, 570, 9, invoice.issue_date)
    text(160, 585, 8, "JATUH TEMPO", True, (0.35, 0.38, 0.43))
    text(160, 570, 9, invoice.due_date)
    commands.append(f"{accent[0]:.3f} {accent[1]:.3f} {accent[2]:.3f} rg 40 530 532 26 re f")
    text(50, 539, 8, "DESKRIPSI", True, (1, 1, 1)); text(365, 539, 8, "QTY", True, (1, 1, 1)); text(420, 539, 8, "HARGA", True, (1, 1, 1)); text(505, 539, 8, "TOTAL", True, (1, 1, 1))
    y = 510
    for item in invoice.items[:12]:
        text(50, y, 8, item.description[:48]); text(370, y, 8, f"{item.quantity:g} {item.unit}"); text(420, y, 8, f"Rp {item.price:,.0f}"); text(505, y, 8, f"Rp {item.total:,.0f}")
        commands.append(f"0.9 0.9 0.9 RG 40 {y - 8} m 572 {y - 8} l S")
        y -= 25
    summary_y = max(170, y - 10)
    text(390, summary_y, 8, "Subtotal", False, (0.35, 0.38, 0.43)); text(490, summary_y, 9, f"Rp {invoice.subtotal:,.0f}", True)
    text(390, summary_y - 20, 8, f"Pajak ({invoice.tax_rate:g}%)", False, (0.35, 0.38, 0.43)); text(490, summary_y - 20, 9, f"Rp {invoice.tax:,.0f}", True)
    commands.append(f"{accent[0]:.3f} {accent[1]:.3f} {accent[2]:.3f} rg 380 {summary_y - 60} 192 30 re f")
    text(390, summary_y - 49, 10, "TOTAL", True, (1, 1, 1)); text(465, summary_y - 49, 11, f"Rp {invoice.total:,.0f}", True, (1, 1, 1))
    text(40, 130, 8, "TERBILANG", True, accent)
    for index, line in enumerate(_wrap(invoice.amount_words, 78)[:2]):
        text(40, 115 - index * 12, 8, line, False, (0.35, 0.38, 0.43))
    text(40, 75, 8, "PEMBAYARAN", True, accent); text(40, 60, 8, f"{company.bank_name} · {company.bank_account}")
    text(320, 75, 8, "CATATAN", True, accent); text(320, 60, 8, invoice.notes[:55], False, (0.35, 0.38, 0.43))
    content = "\n".join(commands).encode("latin-1", "replace")
    jpeg = logo[0] if logo else None
    resources = "<< /Font << /F1 4 0 R /F2 5 0 R >>" + (" /XObject << /Im1 7 0 R >>" if jpeg else "") + " >>"
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources {resources} /Contents 6 0 R >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream",
    ]
    if jpeg and logo:
        _, width, height = logo
        objects.append(f"<< /Type /XObject /Subtype /Image /Width {width} /Height {height} /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode /Length {len(jpeg)} >>\nstream\n".encode() + jpeg + b"\nendstream")
    pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(pdf)); pdf.extend(f"{index} 0 obj\n".encode()); pdf.extend(obj); pdf.extend(b"\nendobj\n")
    xref = len(pdf); pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]: pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(pdf)


def build_invoice_xlsx(invoice: Invoice, company: CompanyProfile, customer: Customer | None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoice"
    accent = company.accent_color.lstrip("#").upper()
    sheet["A1"] = company.name; sheet["A1"].font = Font(size=18, bold=True, color=accent)
    sheet["A2"] = company.address; sheet["A3"] = f"{company.email} | {company.phone}"
    sheet["E1"] = "INVOICE"; sheet["E1"].font = Font(size=20, bold=True, color=accent)
    sheet["E2"] = invoice.number
    sheet["A5"] = "Tagih kepada"; sheet["A5"].font = Font(bold=True)
    sheet["A6"] = invoice.customer; sheet["A7"] = customer.address if customer else ""
    sheet["E5"] = "Tanggal"; sheet["F5"] = invoice.issue_date
    sheet["E6"] = "Jatuh tempo"; sheet["F6"] = invoice.due_date
    headers = ["Deskripsi", "Unit", "Qty", "Harga", "Total"]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(row=10, column=column, value=value); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=accent); cell.alignment = Alignment(horizontal="center")
    for row, item in enumerate(invoice.items, 11):
        sheet.cell(row=row, column=1, value=item.description); sheet.cell(row=row, column=2, value=item.unit); sheet.cell(row=row, column=3, value=item.quantity); sheet.cell(row=row, column=4, value=item.price); sheet.cell(row=row, column=5, value=item.total)
    summary_row = 11 + len(invoice.items) + 1
    sheet.cell(row=summary_row, column=4, value="Subtotal"); sheet.cell(row=summary_row, column=5, value=invoice.subtotal)
    sheet.cell(row=summary_row + 1, column=4, value=f"Pajak ({invoice.tax_rate:g}%)"); sheet.cell(row=summary_row + 1, column=5, value=invoice.tax)
    sheet.cell(row=summary_row + 2, column=4, value="TOTAL").font = Font(bold=True); sheet.cell(row=summary_row + 2, column=5, value=invoice.total).font = Font(bold=True)
    sheet.cell(row=summary_row + 4, column=1, value="Terbilang").font = Font(bold=True); sheet.cell(row=summary_row + 5, column=1, value=invoice.amount_words)
    sheet.cell(row=summary_row + 7, column=1, value=f"Pembayaran: {company.bank_name} · {company.bank_account}")
    for column in "ABCDE": sheet.column_dimensions[column].width = 22
    sheet.column_dimensions["F"].width = 18
    for row in sheet.iter_rows(min_row=11, max_row=summary_row + 2, min_col=4, max_col=5):
        for cell in row: cell.number_format = 'Rp #,##0'
    output = BytesIO(); workbook.save(output); return output.getvalue()