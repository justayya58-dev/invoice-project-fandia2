from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from models.auditpro import TaxDocument, TaxSummary


def build_tax_xlsx(summary: TaxSummary, documents: list[TaxDocument]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Ringkasan"
    summary_sheet["A1"] = "AuditPro AI — Rekap Masa Pajak"
    summary_sheet["A1"].font = Font(size=16, bold=True, color="059669")
    summary_sheet["A2"] = "Masa Pajak"; summary_sheet["B2"] = summary.period
    rows = [
        ("DPP Keluaran", summary.output_dpp), ("PPN Keluaran", summary.output_vat),
        ("DPP Masukan", summary.input_dpp), ("PPN Masukan", summary.input_vat),
        ("PPN Neto", summary.net_vat), ("Siap e-Faktur", summary.ready_count),
        ("Perlu Perbaikan", summary.issue_count),
    ]
    for index, (label, value) in enumerate(rows, 4):
        summary_sheet.cell(row=index, column=1, value=label)
        summary_sheet.cell(row=index, column=2, value=value)
    detail = workbook.create_sheet("Dokumen Pajak")
    headers = ["Arah", "Sumber", "Referensi", "Pihak", "NPWP", "Nomor Faktur", "Tanggal", "Masa", "DPP", "Tarif", "PPN", "Status", "Kelengkapan"]
    for column, value in enumerate(headers, 1):
        cell = detail.cell(row=1, column=column, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="059669")
    for row, item in enumerate(documents, 2):
        values = [item.direction, item.source_type, item.source_reference, item.counterparty, item.npwp, item.tax_invoice_number, item.document_date, item.tax_period, item.dpp, item.tax_rate, item.vat, item.status, item.completeness_score]
        for column, value in enumerate(values, 1):
            detail.cell(row=row, column=column, value=value)
    for column in range(1, len(headers) + 1):
        detail.column_dimensions[chr(64 + column)].width = 18
    for row in detail.iter_rows(min_row=2, min_col=9, max_col=11):
        for cell in row:
            cell.number_format = 'Rp #,##0'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()